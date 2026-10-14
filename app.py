import streamlit as st
import pandas as pd
import json, io, hashlib
from datetime import datetime
from zoneinfo import ZoneInfo
from PIL import Image, ImageOps
from google import genai
from google.genai import types
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

st.set_page_config(page_title="Lector de Remitos", page_icon="📦", layout="wide")

def secreto(nombre):
    return str(st.secrets.get(nombre, "")).strip().strip('"').strip("'")

# ---------- Contraseña (protege tu cuota de IA) ----------
clave = secreto("APP_PASSWORD")
if clave and st.text_input("Contraseña", type="password") != clave:
    st.stop()

api_key = secreto("GEMINI_API_KEY")
if not api_key:
    st.error("Falta GEMINI_API_KEY en Secrets.")
    st.stop()
client = genai.Client(api_key=api_key)

COLS = ["Nº Remito", "Fecha", "Cliente", "Domicilio", "Localidad", "CUIT",
        "Teléfono", "Celular", "Condición", "Detalle de Productos"]
ANCHOS = [20, 14, 35, 32, 28, 18, 12, 16, 35, 55]
CENTRADAS = {0, 1, 5, 6, 7}
CON_AJUSTE = {8, 9}

MODELOS = ["gemini-3.6-flash", "gemini-3.7-flash", "gemini-3.5-flash"]
MAX_ARCHIVOS = 10          # archivos por consulta
MAX_MB = 15                # tamaño total por consulta

PROMPT = """Recibís uno o varios archivos (fotos o PDF) de remitos. Cada archivo puede tener uno o varios remitos.
Devolvé SOLO un JSON: una lista con un objeto por remito, en el mismo orden de los archivos. Cada objeto tiene estas claves:
numero (ej: "00020 - 00002467"), fecha (dd/mm/aaaa), cliente (código y nombre tal como figura, ej: "C305634 - Apellido Nombre"),
domicilio, localidad (tal como figura, con el código postal entre paréntesis), cuit, telefono, celular, condicion,
items (lista de objetos con cantidad, codigo, descripcion).
Copiá el texto exactamente como está impreso. Si un dato no figura, poné "-". No inventes nada."""

def preparar(f):
    data = f.getvalue()
    if getattr(f, "name", "").lower().endswith(".pdf"):
        return types.Part.from_bytes(data=data, mime_type="application/pdf"), len(data)
    img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
    img.thumbnail((1800, 1800))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    jpg = buf.getvalue()
    return types.Part.from_bytes(data=jpg, mime_type="image/jpeg"), len(jpg)

def leer(partes):
    errores = []
    for modelo in MODELOS:
        try:
            r = client.models.generate_content(
                model=modelo,
                contents=partes + [PROMPT],
                config=types.GenerateContentConfig(response_mime_type="application/json", temperature=0),
            )
            break
        except Exception as e:
            errores.append(f"{modelo}: {str(e)[:160]}")
    else:
        raise RuntimeError(" | ".join(errores))
    d = json.loads(r.text)
    if isinstance(d, dict):
        d = [d]
    claves = ["numero", "fecha", "cliente", "domicilio", "localidad",
              "cuit", "telefono", "celular", "condicion"]
    filas = []
    for x in d:
        detalle = "\n".join(
            f"{i.get('cantidad', '1')} x [{i.get('codigo', '')}] {i.get('descripcion', '')}"
            for i in x.get("items", [])
        )
        filas.append([str(x.get(k) or "-") for k in claves] + [detalle or "-"])
    return filas, modelo

def crear_excel(df):
    df = df.fillna("")
    wb = Workbook()
    ws = wb.active
    ws.title = "Remitos Procesados"
    ws.merge_cells("A1:J1")
    ws["A1"] = "Resumen de Remitos Procesados"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
    ws.row_dimensions[1].height = 25
    hoy = datetime.now(ZoneInfo("America/Argentina/Buenos_Aires"))
    ws["A2"] = f"Fecha de procesamiento: {hoy:%d/%m/%Y} | Total Remitos: {len(df)}"
    ws["A2"].font = Font(name="Calibri", size=10, color="595959")
    ws.row_dimensions[4].height = 28
    lado = Side(style="thin", color="D9D9D9")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    for i, titulo in enumerate(COLS):
        c = ws.cell(row=4, column=i + 1, value=titulo)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[chr(65 + i)].width = ANCHOS[i]
    for n, fila in enumerate(df.itertuples(index=False)):
        r = 5 + n
        for i, v in enumerate(fila):
            c = ws.cell(row=r, column=i + 1, value=v)
            c.font = Font(name="Calibri", size=10, color="000000")
            c.alignment = Alignment(horizontal="center" if i in CENTRADAS else "left",
                                    vertical="top", wrap_text=(i in CON_AJUSTE) or None)
            c.border = borde
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F5F9")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

# ---------- Pantalla ----------
st.title("📦 Lector de Remitos a Excel")
if "df" not in st.session_state:
    st.session_state.df = pd.DataFrame(columns=COLS)
    st.session_state.hechos = set()

tab1, tab2 = st.tabs(["📷 Sacar foto", "📁 Subir archivos"])
fotos = []
with tab1:
    c = st.camera_input("Sacá la foto del remito")
    if c:
        fotos.append(c)
with tab2:
    fotos += st.file_uploader("Elegí una o varias imágenes o PDF", type=["jpg", "jpeg", "png", "pdf"],
                              accept_multiple_files=True)

if fotos and st.button("🚀 Procesar", type="primary"):
    lotes, actual, peso = [], [], 0
    for f in fotos:
        h = hashlib.md5(f.getvalue()).hexdigest()
        if h in st.session_state.hechos:
            continue
        try:
            parte, n = preparar(f)
        except Exception as e:
            st.error(f"No se pudo abrir {getattr(f, 'name', 'la foto')}: {e}")
            continue
        if actual and (len(actual) >= MAX_ARCHIVOS or peso + n > MAX_MB * 1024 * 1024):
            lotes.append(actual)
            actual, peso = [], 0
        actual.append((h, parte))
        peso += n
    if actual:
        lotes.append(actual)
    for lote in lotes:
        with st.spinner(f"Leyendo {len(lote)} archivo(s) en una sola consulta..."):
            try:
                filas, modelo = leer([p for _, p in lote])
                for fila in filas:
                    st.session_state.df.loc[len(st.session_state.df)] = fila
                st.session_state.hechos.update(h for h, _ in lote)
                st.caption(f"Leído con {modelo}: {len(lote)} archivo(s), {len(filas)} remito(s).")
                if len(filas) != len(lote):
                    st.warning("La cantidad de remitos no coincide con la de archivos. Revisá la tabla.")
            except Exception as e:
                st.error(f"No se pudo leer: {e}")

if not st.session_state.df.empty:
    st.caption("Revisá los datos: podés corregir cualquier celda antes de descargar.")
    st.session_state.df = st.data_editor(st.session_state.df, use_container_width=True,
                                         hide_index=True, num_rows="dynamic")
    c1, c2 = st.columns(2)
    c1.download_button("📥 Descargar Excel", crear_excel(st.session_state.df),
                       file_name="Remitos_Procesados.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if c2.button("🗑️ Empezar de nuevo"):
        st.session_state.df = pd.DataFrame(columns=COLS)
        st.session_state.hechos = set()
        st.rerun()
